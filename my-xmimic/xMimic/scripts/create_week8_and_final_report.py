"""Generate the week-8 internship report and final production-practice report."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path("/home/eai/my project")
REPO = ROOT / "my-xmimic" / "xMimic"
OUT_WEEK = ROOT / "2023010366_8_黄诗琅.docx"
OUT_FINAL = ROOT / "2023010366_黄诗琅_生产实习最终总结报告.docx"


def set_east_asia(run, name: str) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def configure_document(doc: Document, *, weekly: bool = False) -> None:
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.3 if not weekly else 2.5)
    sec.bottom_margin = Cm(2.3 if not weekly else 2.5)
    sec.left_margin = Cm(2.7 if not weekly else 3.0)
    sec.right_margin = Cm(2.5 if not weekly else 3.0)
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(12)
    pf = style.paragraph_format
    pf.line_spacing = 1.15
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if not weekly:
        add_page_number(sec.footer.paragraphs[0])


def add_body(doc: Document, text: str, *, first_indent: bool = True, after=0) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15
    pf.space_after = Pt(after)
    if first_indent:
        pf.first_line_indent = Pt(24)
    run = p.add_run(text)
    set_east_asia(run, "宋体")
    run.font.size = Pt(12)


def add_h1(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_east_asia(run, "黑体")
    run.font.size = Pt(15)
    run.bold = True


def add_h2(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_east_asia(run, "黑体")
    run.font.size = Pt(13.5)
    run.bold = True


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_east_asia(run, "宋体")
    run.font.size = Pt(10.5)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False if widths else True
    for i, title in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = title
        set_cell_shading(cell, "D9EAF7")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        if widths:
            cell.width = Cm(widths[i])
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                set_east_asia(run, "黑体")
                run.bold = True
                run.font.size = Pt(9.5)
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            if widths:
                cells[i].width = Cm(widths[i])
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(value) < 28 else WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_east_asia(run, "宋体")
                    run.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def create_week8() -> None:
    doc = Document()
    configure_document(doc, weekly=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("0824—0830 实习周报")
    set_east_asia(r, "黑体")
    r.font.size = Pt(18)
    r.bold = True
    info = doc.add_paragraph()
    info.paragraph_format.line_spacing = 1.15
    ir = info.add_run(
        "姓名：黄诗琅\n"
        "实习单位：北京人形机器人创新中心有限公司\n"
        "实习时间：2026 年 8 月 24 日—2026 年 8 月 30 日"
    )
    set_east_asia(ir, "宋体")
    ir.font.size = Pt(12)

    paragraphs = [
        "本周为实习第八周，主要工作是完成上一周复杂复合动作策略的后续精度优化、最终模型筛选、完整回放、ONNX 导出验证和 MuJoCo sim2sim 部署，并系统整理两个月以来形成的代码、动作数据、训练日志、周报和实验结论。此前构建的 complex_choreography_v1 由 FightAndSports、Jumps、PushAndStumble 和 Fight 四段动作及三段平滑过渡组成，共 7319 帧、73.18 秒。本周不再以“训练完成”作为结论，而是继续围绕身体位置、关节姿态、骨盆朝向和水平路径四类 P95 指标做候选 checkpoint 的完整轨迹比较。",
        "针对上一阶段 model_56999 身体位置和关节误差仍略高于严格门槛的问题，我先完成了 4096 环境、10 秒随机相位的 5000 次迭代续训，并对 model_59500 至 model_61998 的六个保存点逐一执行 73.19 秒回放。六个模型均完整执行、无提前终止，其中 model_61000 在水平路径、关节和朝向指标上表现最好，但身体位置 P95 约 0.102 m，仍高于 0.080 m 目标。随后开展身体精度强化、平衡恢复和最终联合精修三轮实验。身体精度阶段一度把身体 P95 降至约 0.078 m，但关节和路径误差回退；平衡恢复阶段又将路径和关节拉回。该过程清楚体现了多目标强化学习中不同奖励之间的权衡，也说明不能只看最后一个 checkpoint 或平均 reward。",
        "最终联合精修使用 4096 个环境追加 2000 次迭代，并在统一配置下评估 model_68500、69000、69500、70000 和 70499。综合完整性和误差后选择 model_69500：关节位置 P95 约 0.829 rad、骨盆朝向 P95 约 0.090 rad，均满足设定门槛，身体位置与水平路径则处于当前训练预算下的最佳折中。随后在名义 seed42 和随机动力学 seed0、1、2 下进行四次完整回放，四次均覆盖全部 7319 帧，没有倒地、越界、NaN 或动作回绕；其中 seed2 严格通过，另外三次主要因身体位置 P95 或水平路径 P95 略超阈值而未通过。因此我将结论明确记录为“稳定执行与部署通过、严格姿态指标部分通过”，没有通过放宽阈值改写实验结果。",
        "工程验证方面，我完成了 checkpoint、ONNX 和视频三条一致性检查。最终策略 Actor 输入为 130 维、Critic 输入为 346 维、输出为 23 个策略关节，跟踪 24 个刚体；ONNX 的权重、前向输出、动作文件哈希、关节和刚体元数据均与 model_69500 一致。名义回放视频共 7319 帧、100 fps、1280×720，完整解码且无黑帧、重复帧或丢帧。部署时将策略安装到 xMIGCS，并针对 rclpy 无法导入的问题绕过存在错误路径钩子的系统 setup.bash，直接配置项目 ROS overlay，最终通过观测维度、关节映射、PD 增益、控制周期和参考动作长度的离线校验。",
        "MuJoCo sim2sim 最终运行 74.18 秒，仿真器和控制器全程存活，策略期间腿、手臂和腰各收到 7396 组控制命令。机器人初始机身高度约 0.965 m，最低高度约 0.575 m，结束高度约 0.972 m，最终向上轴 Z 约 0.9998，最大 XY 半径约 2.336 m，说明机器人完成整套动作后仍保持直立且未越出平台。本周还按照课程手册核对全部周报、中期报告和 my project 下各工程目录，开始形成最终总结报告。通过八周实践，我更加认识到机器人强化学习研发应采用工业工程中的过程质量思想：数据、接口、训练、验收和部署必须形成可追溯证据链；失败模型与未达标指标也应被如实保留。后续将继续完善困难帧采样、多目标奖励平衡和面向真机的安全验证。",
    ]
    for text in paragraphs:
        add_body(doc, text, after=4)
    doc.core_properties.title = "第8周生产实习周报"
    doc.core_properties.author = "黄诗琅"
    doc.save(OUT_WEEK)


def create_final() -> None:
    doc = Document()
    configure_document(doc)

    # Cover
    for _ in range(2):
        doc.add_paragraph()
    for text, size in [
        ("清华大学工业工程系", 18),
        ("《工业工程生产实践》", 22),
        ("生产实习最终总结报告", 26),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(12)
        r = p.add_run(text)
        set_east_asia(r, "黑体")
        r.bold = True
        r.font.size = Pt(size)
    doc.add_paragraph()
    for text in [
        "姓名：黄诗琅",
        "学号：2023010366",
        "实习单位：北京人形机器人创新中心有限公司",
        "实习时间：2026年7月6日—2026年8月30日",
        "校内导师：刘俊驿",
        "2026年8月",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run(text)
        set_east_asia(r, "宋体")
        r.font.size = Pt(14)
    doc.add_page_break()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("生产实习最终总结报告")
    set_east_asia(r, "黑体")
    r.font.size = Pt(18)
    r.bold = True

    add_h1(doc, "一、背景信息")
    add_h2(doc, "1. 实习单位与工作方向")
    add_body(doc, "本人在北京人形机器人创新中心有限公司开展为期八周的生产实习。实习单位聚焦人形机器人关键技术、系统集成与应用验证，我所在的项目工作方向主要涉及人形机器人强化学习、全身动作模仿、仿真评估和策略部署。相关工作既包含 PPO、AMP、Actor-Critic 等算法研究，也包含动作数据处理、Isaac Lab/Isaac Sim 并行仿真、模型导出、ROS 2 控制和 MuJoCo 跨仿真器验证，具有算法研发与工程运作紧密结合的特点。")
    add_body(doc, "本人的实践职责可以概括为：理解并搭建训练环境；梳理机器人关节、观测、奖励和控制接口；执行大规模并行训练；建立完整轨迹、多指标和多随机种子的模型验收方法；将 checkpoint 导出为 ONNX 并接入 xMIGCS、xSIM_MUJOCO；对失败原因进行数据和代码级诊断；最终形成可复现的脚本、配置、日志、视频和报告。为避免涉及企业内部敏感信息，本文仅总结本人在公开项目、公开动作数据和本地仿真环境中完成的技术实践。")

    add_h2(doc, "2. 八周工作安排")
    add_table(doc, ["阶段", "工作重点", "主要成果"], [
        ["第1周", "方向与代码熟悉", "学习VLA/WAM、PPO及Open-X-Humanoid项目结构，明确训练—评估—部署路线。"],
        ["第2周", "TienKung-Lab训练", "配置环境，以1024环境完成walk任务50000次迭代，理解PPO、GAE与AMP。"],
        ["第3周", "结果分析与xMimic切换", "整理回放、导出、sim2sim与Git LFS；完成xMimic代码链路和接口分析。"],
        ["第4周", "dance1长时程优化", "建立状态反馈和分阶段训练，完整回放4/4严格通过，XY P95平均改善77.2%。"],
        ["第5周", "多动作与部署", "完成dance2、weibst训练评估；打通xMIGCS/xSIM_MUJOCO，并构建LAFAN1数据链路。"],
        ["第6周", "六类动作批处理", "完成六类LAFAN1动作训练、评估、ONNX与MuJoCo验证；严格通过16/24。"],
        ["第7周", "复杂长动作生成", "构建73秒复合动作，修复关节顺序错误，完成迁移训练和严格精度续训。"],
        ["第8周", "最终优化与总结", "选出model_69500，完成视频、ONNX、部署和74.18秒sim2sim验证，整理最终报告。"],
    ], widths=[1.5, 3.5, 10.0])

    add_h1(doc, "二、典型一线生产、运作实践与调研总结")
    add_h2(doc, "1. 所参与的研发运作实践")
    add_body(doc, "本次实习所参与的一线实践不是单独运行某个训练脚本，而是一条由动作或任务需求驱动、最终交付可验证策略的研发运作链。前期以TienKung-Lab行走策略为对象，完成依赖安装、AMP-PPO训练、日志解释、策略回放和模型导出；中期转向xMimic全身动作跟踪，围绕dance1、dance2和weibst开展长时程稳定性改善；后期将流程扩展到LAFAN1公开数据的筛选、xGMR重映射、六类动作批量训练，以及73秒复杂复合动作的生成、精修和跨仿真器部署。")
    add_body(doc, "整个流程可划分为七个质量关口：需求与动作选择、数据转换与物理检查、训练入口冒烟、4096环境正式训练、完整轨迹候选筛选、checkpoint/ONNX/视频一致性验证、MuJoCo/ROS闭环验收。任一关口不通过，模型都不进入下一阶段。例如，150秒候选动作虽然完成多轮训练和ONNX导出，但策略在开头过渡后约3.5秒终止，因此没有部署；60秒multiple_actions候选在不同训练路线中均于约20.3秒失败，也被停止继续追加训练。")

    add_h2(doc, "2. 一线研发运作流程")
    add_table(doc, ["流程步骤", "输入/活动", "质量输出"], [
        ["① 数据准备", "BVH/PKL/NPZ筛选、重映射、插值与具名刚体标注", "帧数、频率、哈希、维度、有限值、四元数与边界速度检查"],
        ["② 环境配置", "Isaac Lab、RSL-RL、机器人模型、观测与奖励配置", "64环境冒烟、130/346/23维接口确认"],
        ["③ 并行训练", "1024或4096环境PPO，按时域和难度分阶段训练", "TensorBoard、checkpoint、回合长度与终止分布"],
        ["④ 完整评估", "固定seed与随机动力学条件下逐帧回放", "完成率、XY/身体/关节/朝向P95、平台安全"],
        ["⑤ 模型选择", "比较多个末期checkpoint而非只取最后模型", "综合最佳checkpoint及真实失败项"],
        ["⑥ 导出验证", "导出ONNX、录制视频并校验元数据与前向输出", "权重一致性、视频完整性、可追溯哈希"],
        ["⑦ sim2sim", "xMIGCS控制器、ROS 2通信与MuJoCo物理仿真", "命令/状态频率、机身高度、姿态、边界和异常检查"],
    ], widths=[2.2, 6.0, 7.0])
    add_caption(doc, "图表1  人形机器人策略训练与部署的七阶段质量控制流程")

    add_h2(doc, "3. 实践问题与运作经验")
    add_body(doc, "第一类问题是环境与依赖不一致。实习中先后遇到自定义rsl_rl导入、Isaac Sim必须由AppLauncher初始化、GPU驱动与用户空间库不匹配、ROS 2的rclpy搜索路径以及系统setup.bash包含错误路径钩子等问题。我逐步采用独立Conda环境、可编辑安装、完整解释器路径、仓库内ROS overlay和启动前自检，减少“在某个终端能运行、换一个终端失效”的隐性差异。")
    add_body(doc, "第二类问题是数据和接口错配。xMimic训练端包含23个策略关节，MuJoCo部署端包含29个电机；不同策略又存在124、130和133维观测，不能只修改YAML中的num_obs。动作文件还涉及body_names、Isaac Articulation关节顺序和qpos转换顺序。本次73秒动作初版正是因为两套23维关节顺序混用，往返偏差达到约3.08 rad。修复后最大偏差降至约1.6×10⁻⁴ rad，说明显式名称映射、元数据和往返测试比肉眼预览更可靠。")
    add_body(doc, "第三类问题是训练指标与真实使用目标不一致。平均reward上升、timeout比例增加或ONNX能够生成，并不等于完整动作可用。dance2在Isaac Lab短视频中表现正常，却在MuJoCo约一分钟后倒地；某些末期checkpoint也会比早期模型明显退化。为此我建立完整帧序列、多checkpoint、多随机种子和跨仿真器的验收协议，并将“完整性、安全性、精度、部署可用性”分别统计。这种做法本质上是把工业工程中的过程能力、质量关口和追溯思想应用到算法研发。")

    add_h1(doc, "三、专题项目改善总结")
    add_h2(doc, "1. 项目背景、范围、目标与计划")
    add_body(doc, "个人专题项目为“面向长时程复杂动作的人形机器人全身模仿策略优化与sim2sim验证”。项目以公开动作数据和DexEVT2机器人为对象，覆盖动作筛选与转换、Isaac Lab强化学习、观测和奖励设计、完整轨迹评估、checkpoint选择、ONNX导出、xMIGCS控制和MuJoCo验证。目标不只是提高训练reward，而是使策略在完整动作时域内连续执行，位置和姿态误差可量化，在扰动下保持稳定，并保证训练端与部署端接口一致。")
    add_body(doc, "计划采用逐步扩展方式：先通过TienKung walk掌握PPO/AMP闭环；再以dance1定位长时程漂移并建立状态反馈；把方法迁移至dance2、weibst和六类LAFAN1动作；最后生成包含拳击、跳跃、受推恢复和连续格斗的73秒复合动作，验证从数据生产到MuJoCo部署的完整能力。")

    add_h2(doc, "2. 问题的结构化描述")
    add_table(doc, ["要素", "具体内容"], [
        ["决策变量", "episode时域、起始相位采样、Actor观测项、奖励结构与权重、随机化强度、迁移初始化和checkpoint。"],
        ["评价指标", "动作完成率；XY、身体位置、关节姿态和骨盆朝向P95/最大值；终止原因；平台边界；MuJoCo机身高度和姿态。"],
        ["约束条件", "100 Hz参考动作与控制时序；23维策略动作；130维状态反馈观测；部署端29电机映射；8 m平台安全边界。"],
        ["输入数据", "TienKung任务配置、dance/weibst NPZ、LAFAN1 BVH、xGMR PKL、具名刚体运动数据、训练日志和回放telemetry。"],
        ["质量目标", "完整帧序列、无NaN/回绕/提前终止；名义和随机条件稳定；checkpoint、ONNX、视频和部署配置可追溯。"],
    ], widths=[2.6, 12.6])

    add_h2(doc, "3. 行业实践与技术调研")
    add_body(doc, "算法层面，我结合TienKung-Lab和xMimic代码学习PPO、GAE、AMP和非对称Actor-Critic。PPO通过限制新旧策略更新幅度保持训练稳定；GAE利用价值函数自举估计优势；AMP通过判别器学习专家运动风格；xMimic则把参考动作的关节、刚体与根节点状态作为条件，以运动误差奖励训练全身跟踪策略。工程层面，我阅读并实践Isaac Lab/RSL-RL训练接口、xGMR动作重映射、ONNX模型交换、ROS 2控制和MuJoCo仿真，理解了训练数据、策略网络、控制器和机器人模型之间必须共享明确的维度、顺序、频率和坐标约定。")
    add_body(doc, "工业工程视角下，该项目可视为小批量、多品种的算法生产流程：每条动作是一种“产品”，数据转换和训练是加工过程，完整回放和sim2sim是检验工序，checkpoint与ONNX是交付件。传统只看训练reward的做法相当于只监控设备内部参数而不做出厂检验。本项目通过标准化输入检查、阶段门、量化验收和异常隔离，提高了实验复现性，也防止不合格模型进入部署端。")

    add_h2(doc, "4. 解决方案形成与实施")
    add_body(doc, "（1）建立数据质量前置检查。对NPZ的键、shape、帧率、数值有限性和四元数归一化进行检查；使用body_names按名称解析刚体；对BVH→PKL→NPZ链路记录源片段和哈希；对qpos与训练关节顺序做往返验证。对于错误数据不覆盖删除，而是隔离并标明REJECTED，避免后续误选。")
    add_body(doc, "（2）改进长时程训练。dance1原始episode仅5秒，无法学习39.72秒轨迹后段的累计漂移。我采用5→10→20→完整时域的课程训练，按episode长度限制合法起始相位，避免参考动作越界回绕；同时在Actor中加入参考骨盆相对位置和基座线速度，使观测由124维扩展至130维，并增加全局XY位置、速度和关节保真奖励。后续动作根据长度采用完整时域或随机相位训练。")
    add_body(doc, "（3）建立多目标checkpoint选择。每个训练阶段保留固定间隔的checkpoint，对末期多个候选执行完整轨迹回放。筛选时先排除提前终止、NaN、动作回绕和越界，再综合XY、身体、关节和朝向P95，而不是按最后迭代或单一reward决定。对复杂动作还进行身体精度、平衡恢复和联合精修，观察奖励权重改变造成的指标权衡，并在收益降低时停止继续训练。")
    add_body(doc, "（4）打通部署证据链。对选定checkpoint检查Actor 130维、Critic 346维和23维动作结构；导出ONNX后核对权重、前向输出、动作哈希、关节/刚体元数据和action scale；验证7319帧视频可完整解码；再生成BeyondMimic YAML并检查23个策略关节到29个MuJoCo电机的映射、PD增益和控制周期。最后通过ROS 2将策略节点与MuJoCo连接，记录电机命令、状态和机身telemetry。")

    add_h2(doc, "5. 主要结果与改善效果")
    add_table(doc, ["对象", "规模/时长", "结果与改进"], [
        ["TienKung walk", "1024环境，50000迭代", "完成训练、回放与导出闭环，建立PPO/GAE/AMP代码级理解。"],
        ["dance1", "3972帧，39.72 s", "四工况4/4严格通过；XY P95平均从0.8111 m降至0.1848 m，改善约77.2%。"],
        ["dance2", "10299帧，102.99 s", "完整时域续训后四工况均完成；seed1 XY P95由4.763 m降至0.213 m。"],
        ["weibst_hard", "1327帧，50000迭代", "四工况完整、安全，综合跟踪分数改善14.17%；严格姿态指标仍部分未达标。"],
        ["LAFAN1六类", "6动作，24工况", "24/24完整安全、16/24严格通过；六策略ONNX、离线部署和MuJoCo全部通过。"],
        ["150秒候选", "约157.6 s，多阶段训练", "发现开头IK/截帧异常并重建数据；恢复路线仍在约4 s终止，未部署。"],
        ["60秒多动作", "约67.5 s，多路线", "不同模型均在约20.3 s重复失败，判断为参考动作高风险段，停止盲目续训。"],
        ["73秒复合动作", "7319帧，4096环境", "model_69500四工况均完整；严格1/4通过；74.18 s MuJoCo不倒地、不越界。"],
    ], widths=[2.5, 3.1, 9.6])

    training_img = REPO / "artifacts/2026-08-19_complex_choreography_v1/final_balance_training_analysis/training_overview.png"
    if training_img.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(training_img), width=Cm(13.8))
        add_caption(doc, "图1  complex_choreography_v1最终联合精修训练概览")
    telemetry_img = REPO / "artifacts/2026-08-19_complex_choreography_v1/sim2sim_validation/telemetry.png"
    if telemetry_img.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(telemetry_img), width=Cm(13.8))
        add_caption(doc, "图2  73秒复合动作MuJoCo sim2sim状态记录")

    add_body(doc, "最终73秒复合动作由FightAndSports、Jumps、PushAndStumble和Fight四段及三个20帧平滑桥接组成，水平范围约2.40 m×2.20 m。最终model_69500在四次Isaac Lab回放中均完成7319帧，无终止、回绕、NaN或越界；严格指标为1/4通过，未通过工况主要卡在身体位置和少量XY漂移，因此结论为“稳定执行/部署通过、严格姿态部分通过”。MuJoCo运行74.18秒时，机身最低高度0.575 m、结束高度0.972 m、最终向上轴Z为0.9998、最大XY半径2.336 m，仿真器和控制器全程存活。")

    add_h2(doc, "6. 工程文件与成果资产总结")
    add_body(doc, "截至2026年8月28日，my project目录形成了由学习仓库、训练仓库、数据源、部署端、仿真端和课程文档组成的完整资产。下表按顶层目录归纳其用途；模型、日志和大规模数据通过Git LFS或本地归档管理，报告中不逐项罗列数万个中间文件。")
    add_table(doc, ["目录/文件组", "作用与成果"], [
        ["my-tienkung-robot（约23 GB）", "TienKung-Lab学习、walk训练日志、模型、PPO/GAE/AMP学习笔记与早期sim2sim实践。"],
        ["xmimic-old（约498 MB）", "xMimic原始/旧版基线，用于对照README、动作入口和历史训练命令。"],
        ["my-xmimic（约24 GB）", "主要研发仓库；包含动作数据、训练日志、checkpoint、评估脚本、视频、ONNX和全流程artifacts。"],
        ["motion_sources（约2.2 GB）", "LAFAN1公开BVH与xGMR重映射工具，支撑六类动作和长动作数据生产。"],
        ["Deploy_Tienkung及fresh副本", "部署代码与README基线，用于核对控制节点、依赖和官方结构，避免误改后无基线。"],
        ["xSIM_MUJOCO（约50 MB）", "MuJoCo物理仿真器及evt2机器人模型，负责接收控制命令并发布状态。"],
        ["xmigcs_test（约830 MB）", "实际xMIGCS策略控制与验证仓库，保存多动作BeyondMimic YAML、ONNX及ROS overlay。"],
        ["周报/中期报告/课程手册", "第1—8周过程记录、中期总结与最终报告依据，保证课程过程和技术成果一致。"],
    ], widths=[4.4, 10.8])

    add_h2(doc, "7. 局限与后续改进")
    add_body(doc, "第一，部分策略虽然完整执行和sim2sim通过，但身体关键点P95未满足严格门槛，说明稳定性与动作保真度仍需分别优化。第二，现有策略以单动作单策略为主，多动作统一策略和未见动作泛化尚未验证。第三，仿真验证不能替代真机安全测试，真实执行还需设置力矩、速度、姿态和区域保护，并遵循企业审批流程。后续可采用困难帧分层采样、动作分段课程、奖励尺度自动平衡、历史观测或蒸馏方法改善高难段，同时用质量、摩擦、执行器延迟和push单因素实验识别sim2sim敏感参数。")

    add_h1(doc, "四、生产实习总结")
    add_h2(doc, "1. 客观总结")
    add_body(doc, "八周实习完成了从基础学习到独立构建端到端实验流程的过渡。我完成1024环境TienKung walk正式训练，掌握PPO、GAE和AMP；完成xMimic环境搭建、核心代码与接口分析；围绕dance1、dance2、weibst和LAFAN1六类动作开展4096环境训练、长时程优化和严格评估；生成并验证73秒复合动作；打通checkpoint、ONNX、xMIGCS、ROS 2和MuJoCo；同时建立数据校验、候选筛选、视频/模型一致性、sim2sim telemetry和Git/Git LFS归档方法。")

    add_h2(doc, "2. 工业工程认识")
    add_body(doc, "本次实践使我认识到，算法研发同样需要过程设计与质量管理。动作数据是原材料，训练配置和计算资源是工艺条件，checkpoint是半成品，完整回放、ONNX验证和sim2sim是检验工序。若只追求训练轮数或reward，相当于忽略顾客真正关心的完整性、安全性和可部署性。通过流程分解、指标体系、阶段门、异常隔离、版本控制和可视化，我把零散调试转化为可以复现和审核的运作过程。这也是工业工程方法在智能机器人研发中的具体体现。")

    add_h2(doc, "3. 个人收获与体会")
    add_body(doc, "技术上，我建立了从强化学习理论到机器人控制代码的对应关系，能够分析观测、奖励、终止条件、策略网络、关节映射和仿真通信；工程上，提升了Linux、Conda、Git LFS、TensorBoard、ONNX、ROS 2和MuJoCo的综合使用能力；问题解决上，形成“复现现象—划分模块—提出假设—做对照实验—保留证据—验证修复”的习惯。尤其重要的是，我学会如实报告失败：150秒动作和60秒候选没有通过，complex动作严格指标只有1/4通过，这些结果同样揭示了数据质量和多目标权衡，对后续改进有价值。")
    add_body(doc, "在纪律、安全与职业素养方面，我遵守实习单位和学校的保密、知识产权与考勤要求，仅使用获准的公开数据和本地仿真材料总结成果；对高风险机器人动作坚持先仿真、再部署验证，不把未经完整验收的模型用于真机。与导师和同事交流时，我逐步学会用流程、数据和证据说明问题，而不是只描述主观观感。")

    add_h2(doc, "4. 后续展望")
    add_body(doc, "后续我将继续围绕复杂动作的数据质量、困难阶段训练和sim2real开展研究：一是完善自动化动作审计与困难帧采样；二是研究多目标奖励的动态平衡，减少身体精度与全局路径、关节保真之间的冲突；三是扩展多动作统一策略及更长时间动作；四是在满足审批和安全条件下，逐级开展硬件在环与真机验证。希望把本次实习形成的标准化训练—验收—部署流程沉淀为可复用工具，为人形机器人运动技能开发提供稳定基础。")

    add_h1(doc, "参考资料")
    refs = [
        "[1] Open-X-Humanoid. TienKung-Lab 项目代码与说明。",
        "[2] Open-X-Humanoid. xMimic 项目代码与README。",
        "[3] Open-X-Humanoid. xGMR 动作重映射工具。",
        "[4] NVIDIA. Isaac Lab / Isaac Sim 文档与示例。",
        "[5] RSL-RL. On-policy reinforcement learning implementation。",
        "[6] LAFAN1 Motion Capture Dataset。",
        "[7] MuJoCo Physics Simulator。",
        "[8] 《2026年工业工程生产实践课程手册》。",
    ]
    for text in refs:
        add_body(doc, text, first_indent=False)

    doc.core_properties.title = "生产实习最终总结报告"
    doc.core_properties.subject = "人形机器人强化学习、动作模仿与sim2sim验证"
    doc.core_properties.author = "黄诗琅"
    doc.core_properties.keywords = "工业工程生产实践; 人形机器人; 强化学习; xMimic; sim2sim"
    doc.save(OUT_FINAL)


if __name__ == "__main__":
    create_week8()
    create_final()
    print(OUT_WEEK)
    print(OUT_FINAL)
